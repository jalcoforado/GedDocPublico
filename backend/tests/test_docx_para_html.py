"""DOCX → HTML da sincronização de minuta com o Google Docs (backlog 2.4 (b)).

Até 2026-09-29 só o texto dos parágrafos voltava: negrito, listas, títulos,
alinhamento e tabelas se perdiam — para documento oficial, isso tornava a
sincronização quase inútil. A conversão agora é uma função pura, testada aqui
sem banco e sem Google: os DOCX são montados com python-docx.

Regra que não pode quebrar: texto SEM formatação sai idêntico ao formato antigo
(`<p>…</p>`). Se mudasse, a primeira sincronização depois do deploy criaria
versão nova em toda minuta sem ninguém ter mexido no Google.
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.docx_para_html import docx_para_html
from app.services.html_sanitizer import sanitizar_html


def _bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _html(doc) -> str:
    """Como a sincronização usa: converte e sanitiza."""
    return sanitizar_html(docx_para_html(_bytes(doc)))


def test_texto_sem_formatacao_sai_identico_ao_formato_antigo():
    d = Document()
    for t in ("Primeira linha", "", "   ", "Segunda linha"):
        d.add_paragraph(t)
    assert _html(d) == "<p>Primeira linha</p><p>Segunda linha</p>"


def test_texto_continua_escapado():
    d = Document()
    d.add_paragraph("Prezado <nome>, valor > 10 & mais")
    d.add_paragraph("<script>alert(1)</script>")
    h = _html(d)
    assert "Prezado &lt;nome&gt;, valor &gt; 10 &amp; mais" in h
    assert "<script" not in h and "&lt;script&gt;" in h


def test_negrito_italico_sublinhado_tachado():
    d = Document()
    p = d.add_paragraph("Normal ")
    p.add_run("negrito").bold = True
    p.add_run(" e ")
    p.add_run("itálico").italic = True
    p.add_run(" e ")
    p.add_run("sublinhado").underline = True
    p.add_run(" e ")
    r = p.add_run("tachado")
    r.font.strike = True
    assert _html(d) == (
        "<p>Normal <strong>negrito</strong> e <em>itálico</em> e "
        "<u>sublinhado</u> e <s>tachado</s></p>"
    )


def test_runs_vizinhos_com_mesma_formatacao_viram_uma_tag_so():
    """O Google quebra um trecho em vários runs; sem fundir, sairia
    `<strong>a</strong><strong>b</strong>`."""
    d = Document()
    p = d.add_paragraph()
    p.add_run("Art. ").bold = True
    p.add_run("1º").bold = True
    p.add_run(" — texto")
    assert _html(d) == "<p><strong>Art. 1º</strong> — texto</p>"


def test_negrito_e_italico_juntos():
    d = Document()
    r = d.add_paragraph().add_run("ênfase")
    r.bold = True
    r.italic = True
    assert _html(d) == "<p><strong><em>ênfase</em></strong></p>"


def test_titulos_viram_os_niveis_que_o_editor_usa():
    """O editor da plataforma (TipTap) só tem títulos 2 e 3."""
    d = Document()
    d.add_heading("Título do documento", level=0)
    d.add_heading("Seção", level=1)
    d.add_heading("Subseção", level=2)
    d.add_heading("Sub-subseção", level=3)
    assert _html(d) == (
        "<h2>Título do documento</h2><h2>Seção</h2><h3>Subseção</h3><h3>Sub-subseção</h3>"
    )


def test_alinhamento_vira_text_align():
    d = Document()
    for texto, al in (("centro", WD_ALIGN_PARAGRAPH.CENTER), ("direita", WD_ALIGN_PARAGRAPH.RIGHT),
                      ("justificado", WD_ALIGN_PARAGRAPH.JUSTIFY), ("esquerda", WD_ALIGN_PARAGRAPH.LEFT)):
        d.add_paragraph(texto).alignment = al
    h = d.add_heading("Seção centrada", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    assert _html(d) == (
        '<p style="text-align: center;">centro</p>'
        '<p style="text-align: right;">direita</p>'
        '<p style="text-align: justify;">justificado</p>'
        "<p>esquerda</p>"
        '<h2 style="text-align: center;">Seção centrada</h2>'
    )


def test_listas_por_estilo_do_word():
    d = Document()
    d.add_paragraph("Considerando:")
    d.add_paragraph("primeiro", style="List Bullet")
    d.add_paragraph("segundo", style="List Bullet")
    d.add_paragraph("Resolve:")
    d.add_paragraph("item um", style="List Number")
    d.add_paragraph("item dois", style="List Number")
    assert _html(d) == (
        "<p>Considerando:</p><ul><li>primeiro</li><li>segundo</li></ul>"
        "<p>Resolve:</p><ol><li>item um</li><li>item dois</li></ol>"
    )


def _lista_direta(d, fmt: str) -> int:
    """Numeração marcada no PRÓPRIO parágrafo, como o Google Docs exporta."""
    numbering = d.part.numbering_part.element
    abs_id = 900 + len(numbering.findall(qn("w:abstractNum")))
    absn = OxmlElement("w:abstractNum")
    absn.set(qn("w:abstractNumId"), str(abs_id))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    nf = OxmlElement("w:numFmt")
    nf.set(qn("w:val"), fmt)
    lvl.append(nf)
    absn.append(lvl)
    numbering.insert(0, absn)
    num_id = 900 + len(numbering.findall(qn("w:num")))
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abs_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def _item(d, texto: str, num_id: int):
    p = d.add_paragraph(texto)
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(nid)
    ppr.append(numpr)
    return p


def test_listas_no_formato_do_google_docs():
    d = Document()
    marcador = _lista_direta(d, "bullet")
    numerada = _lista_direta(d, "decimal")
    _item(d, "a", marcador)
    _item(d, "b", marcador)
    _item(d, "um", numerada)
    _item(d, "dois", numerada)
    assert _html(d) == "<ul><li>a</li><li>b</li></ul><ol><li>um</li><li>dois</li></ol>"


def test_item_de_lista_mantem_formatacao():
    d = Document()
    p = d.add_paragraph(style="List Bullet")
    p.add_run("Prazo: ").bold = True
    p.add_run("30 dias")
    assert _html(d) == "<ul><li><strong>Prazo: </strong>30 dias</li></ul>"


def test_tabela_na_ordem_do_documento():
    d = Document()
    d.add_paragraph("Antes")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Item"
    t.cell(0, 1).text = "Valor"
    t.cell(1, 0).text = "Diárias <2>"
    t.cell(1, 1).paragraphs[0].add_run("R$ 100").bold = True
    d.add_paragraph("Depois")
    assert _html(d) == (
        "<p>Antes</p>"
        "<table><tbody>"
        "<tr><td>Item</td><td>Valor</td></tr>"
        "<tr><td>Diárias &lt;2&gt;</td><td><strong>R$ 100</strong></td></tr>"
        "</tbody></table>"
        "<p>Depois</p>"
    )


def test_quebra_de_linha_dentro_do_paragrafo():
    d = Document()
    r = d.add_paragraph().add_run("linha 1")
    r.add_break()
    r.add_text("linha 2")
    assert _html(d) == "<p>linha 1<br>linha 2</p>"


def test_documento_vazio_vira_string_vazia():
    assert _html(Document()) == ""
