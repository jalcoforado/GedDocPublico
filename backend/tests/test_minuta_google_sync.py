import io
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document

from app.models.minuta import Minuta
from app.services import minutas as svc
from app.services.google_docs_service import GoogleDocsService


def _create_fake_docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Este é um parágrafo do Google Docs.")
    doc.add_paragraph("E este é outro parágrafo.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


@pytest.mark.asyncio
async def test_sincronizar_google_doc_para_minuta():
    """Testa a extração e sincronização de bytes DOCX sem tocar no DB real usando dublês."""
    
    # Criar um mock para a sessão do DB
    db_mock = AsyncMock()

    # Criar um mock para a Minuta
    minuta_mock = Minuta(
        id=1,
        tenant_id=1,
        id_processo=10,
        titulo="Minuta Teste",
        status="rascunho",
        google_doc_id="fake_doc_id",
        versao=1,
        corpo_html="",
    )

    # Patch nas funções dependentes
    with patch("app.services.minutas.obter_minuta", new_callable=AsyncMock) as mock_obter, \
         patch.object(GoogleDocsService, "obter_credentials_usuario", new_callable=AsyncMock) as mock_creds, \
         patch.object(GoogleDocsService, "sincronizar_google_doc", new_callable=AsyncMock) as mock_sync, \
         patch("app.services.minutas.audit_log", new_callable=AsyncMock) as mock_audit:
        
        mock_obter.return_value = minuta_mock
        mock_creds.return_value = "fake_cred"
        mock_sync.return_value = _create_fake_docx_bytes()

        m_atualizada = await svc.sincronizar_google_doc_para_minuta(
            db_mock,
            tenant_id=1,
            minuta_id=1,
            usuario_id=99,
        )

        assert "<p>Este é um parágrafo do Google Docs.</p>" in m_atualizada.corpo_html
        assert "<p>E este é outro parágrafo.</p>" in m_atualizada.corpo_html
        assert m_atualizada.versao == 2
        
        # O mock sync deve ter sido chamado com fake_doc_id
        mock_sync.assert_called_once_with(db_mock, cred="fake_cred", google_doc_id="fake_doc_id")
